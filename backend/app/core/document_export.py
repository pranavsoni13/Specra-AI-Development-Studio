"""Downloadable SRS, SDD, test report and requirements traceability exports."""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _text(value: Any, default: str = "Not specified") -> str:
    if value is None:
        return default
    if isinstance(value, str):
        cleaned = value.strip()
        return cleaned or default
    return str(value)


def _items(value: Any) -> list:
    return value if isinstance(value, list) else []


def _section_list(title: str, values: Any) -> str:
    items = [str(item).strip() for item in _items(values) if str(item).strip()]
    if not items:
        return f"## {title}\n\nNone recorded.\n"
    return f"## {title}\n\n" + "\n".join(f"- {item}" for item in items) + "\n"


def _requirement_rows(requirements: dict | None) -> list[dict]:
    requirements = requirements or {}
    return [
        {"requirement_id": f"FR-{index:03d}", "text": str(text).strip()}
        for index, text in enumerate(
            [
                item for item in _items(requirements.get("functional_requirements"))
                if isinstance(item, str) and item.strip()
            ],
            start=1,
        )
    ]


def render_srs(
    project: dict,
    artifacts: dict,
    versions: dict,
) -> str:
    requirements = artifacts.get("requirements") or {}
    rows = _requirement_rows(requirements)
    lines = [
        f"# Software Requirements Specification (SRS): {_text(requirements.get('project_name'), project.get('project_name') or 'Specra Project')}",
        "",
        f"- Project ID: {project.get('project_id', 'unknown')}",
        f"- Document generated: {_now()}",
        "- Document status: Draft generated from the latest saved requirements artifact",
        f"- Requirements artifact version: {versions.get('requirements', 0)}",
        "",
        "## 1. Purpose and Scope",
        "",
        _text(requirements.get("objective"), project.get("idea", "Purpose not recorded.")),
        "",
        "## 2. Original Project Brief",
        "",
        _text(project.get("idea")),
        "",
        "## 3. Stakeholders and User Roles",
        "",
    ]
    roles = _items(requirements.get("user_roles"))
    lines.extend([*(f"- {role}" for role in roles), *(["No user roles recorded."] if not roles else [])])
    lines.extend(["", "## 4. Functional Requirements", ""])
    if rows:
        lines.extend(
            f"- **{row['requirement_id']}** — {row['text']}"
            for row in rows
        )
    else:
        lines.append("No functional requirements are available in the saved artifact.")
    lines.extend(["", _section_list("5. Non-Functional Requirements", requirements.get("non_functional_requirements")).rstrip()])
    lines.extend(["", _section_list("6. Constraints", requirements.get("constraints")).rstrip()])
    lines.extend(["", _section_list("7. Assumptions", requirements.get("assumptions")).rstrip()])
    lines.extend(["", _section_list("8. Risks", requirements.get("risks")).rstrip()])
    lines.extend(["", _section_list("9. Clarifications Required", requirements.get("clarifications_needed")).rstrip()])
    lines.extend([
        "",
        "## 10. Traceability and Limitations",
        "",
        "Functional IDs are assigned in the saved requirement order and are intended to match Specra's FR-001 style traceability.",
        "This draft reflects the latest saved requirements artifact; confirm the content with stakeholders before treating it as an approved specification.",
        "",
    ])
    return "\n".join(lines)


def render_sdd(
    project: dict,
    artifacts: dict,
    versions: dict,
) -> str:
    requirements = artifacts.get("requirements") or {}
    plan = artifacts.get("project_plan") or {}
    architecture = artifacts.get("architecture") or {}
    name = _text(architecture.get("project_name"), project.get("project_name") or "Specra Project")
    lines = [
        f"# Software Design Description (SDD): {name}",
        "",
        f"- Project ID: {project.get('project_id', 'unknown')}",
        f"- Document generated: {_now()}",
        "- Document status: Draft generated from the latest saved plan and architecture artifacts",
        f"- Project plan version: {versions.get('project_plan', 0)}",
        f"- Architecture version: {versions.get('architecture', 0)}",
        "",
        "## 1. Design Overview",
        "",
        f"- Architecture style: {_text(architecture.get('architecture_style'))}",
        f"- Frontend: {_text(architecture.get('frontend'))}",
        f"- Backend: {_text(architecture.get('backend'))}",
        f"- Database: {_text(architecture.get('database'))}",
        f"- Authentication: {_text(architecture.get('authentication'))}",
        "",
        "## 2. Component Design",
        "",
    ]
    components = _items(architecture.get("components"))
    if components:
        for component in components:
            lines.append(f"### {_text(component.get('name'))}")
            lines.append("")
            lines.append(_text(component.get("responsibility")))
            refs = _items(component.get("related_requirements"))
            lines.append("")
            lines.append("Related requirements: " + (", ".join(refs) if refs else "Not mapped"))
            lines.append("")
    else:
        lines.append("No architecture components are available in the saved artifact.")
        lines.append("")
    lines.extend(["## 3. API Design", ""])
    endpoints = _items(architecture.get("api_endpoints"))
    if endpoints:
        for endpoint in endpoints:
            lines.append(f"- **{_text(endpoint.get('name'))}** — {_text(endpoint.get('purpose'))} (requirements: {', '.join(_items(endpoint.get('related_requirements'))) or 'not mapped'})")
    else:
        lines.append("No API capabilities are recorded.")
    lines.extend(["", "## 4. Data Flow", ""])
    flows = _items(architecture.get("data_flow"))
    lines.extend([*(f"- {item}" for item in flows), *(["No data-flow description recorded."] if not flows else [])])
    lines.extend(["", "## 5. Architecture Decisions", ""])
    decisions = _items(architecture.get("architecture_decisions"))
    if decisions:
        for index, decision in enumerate(decisions, start=1):
            lines.extend([
                f"### Decision {index}: {_text(decision.get('decision'))}",
                "",
                _text(decision.get("rationale")),
                "",
                "Related requirements: " + (", ".join(_items(decision.get("related_requirements"))) or "Not mapped"),
                "",
            ])
    else:
        lines.extend(["No architecture decisions are recorded.", ""])
    lines.extend(["## 6. Implementation Plan", ""])
    milestones = _items(plan.get("milestones"))
    if milestones:
        for index, milestone in enumerate(milestones, start=1):
            lines.extend([
                f"### {index}. {_text(milestone.get('name'))}",
                "",
                _text(milestone.get("description")),
                "",
                "Related requirements: " + (", ".join(_items(milestone.get("requirement_refs"))) or "Not mapped"),
                "",
                "Tasks:",
            ])
            tasks = _items(milestone.get("tasks"))
            if tasks:
                lines.extend(
                    f"- {_text(task.get('description'))} (requirements: {', '.join(_items(task.get('requirement_refs'))) or 'not mapped'})"
                    for task in tasks
                )
            else:
                lines.append("- No tasks recorded.")
            lines.append("")
    else:
        lines.append("No implementation milestones are recorded.")
    lines.extend(["", _section_list("7. External Integrations", architecture.get("external_integrations")).rstrip()])
    lines.extend(["", _section_list("8. Architectural Assumptions", architecture.get("architectural_assumptions")).rstrip()])
    lines.extend(["", _section_list("9. Related Requirements", requirements.get("functional_requirements")).rstrip()])
    lines.extend([
        "",
        "## 10. Design Status and Limitations",
        "",
        f"Technical review version: {versions.get('technical_review', 0)}.",
        "This is a draft design document. Validate security, deployment, operational constraints, and architecture trade-offs with the engineering team before implementation.",
        "",
    ])
    return "\n".join(lines)


def render_test_report(
    project: dict,
    suite: dict | None,
) -> str:
    project_name = _text(
        (suite or {}).get("project_name"),
        project.get("project_name") or "Specra Project",
    )
    cases = _items((suite or {}).get("cases"))
    counts = {
        "passed": sum(1 for case in cases if case.get("status") == "passed"),
        "failed": sum(1 for case in cases if case.get("status") == "failed"),
        "blocked": sum(1 for case in cases if case.get("status") == "blocked"),
        "not_run": sum(1 for case in cases if case.get("status", "not_run") == "not_run"),
    }
    executed = counts["passed"] + counts["failed"] + counts["blocked"]
    lines = [
        f"# Test Report: {project_name}",
        "",
        f"- Project ID: {project.get('project_id', 'unknown')}",
        f"- Report generated: {_now()}",
        f"- Test suite version: {(suite or {}).get('suite_version', 'Not generated')}",
        f"- Test cases: {len(cases)}",
        f"- Cases marked as executed: {executed}",
        f"- Passed: {counts['passed']}",
        f"- Failed: {counts['failed']}",
        f"- Blocked: {counts['blocked']}",
        f"- Not run: {counts['not_run']}",
        "",
        "## Status and Evidence",
        "",
        "This report distinguishes AI-proposed test design from test execution. A case is counted as executed only after an authorized user records a result. Specra does not claim that generated tests passed unless actual execution evidence is recorded.",
        "",
    ]
    if not suite:
        lines.extend([
            "No test suite has been generated for this project yet. Generate and review test cases in the Testing workspace, then record execution results to populate this report.",
            "",
        ])
    elif cases:
        for case in cases:
            lines.extend([
                f"## {case.get('test_id', 'TC')} — {_text(case.get('title'))}",
                "",
                f"- Requirement: {case.get('requirement_id', 'Not mapped')} — {_text(case.get('requirement_text'))}",
                f"- Type: {_text(case.get('test_type'))}",
                f"- Priority: {_text(case.get('priority'))}",
                f"- Status: {case.get('status', 'not_run')}",
                f"- Generation source: {case.get('generation_source', 'unknown')}",
                f"- Executed at: {case.get('executed_at') or 'Not recorded'}",
                f"- Executed by: {case.get('executed_by') or 'Not recorded'}",
                f"- Environment: {case.get('environment') or 'Not recorded'}",
                "",
                "Objective: " + _text(case.get("objective")),
                "",
                "### Preconditions",
                "",
                *([f"- {item}" for item in _items(case.get("preconditions"))] or ["- None recorded."]),
                "",
                "### Steps",
                "",
                *([f"{index}. {step}" for index, step in enumerate(_items(case.get("steps")), start=1)] or ["No steps recorded."]),
                "",
                "### Expected Result",
                "",
                _text(case.get("expected_result")),
                "",
                "### Execution Notes",
                "",
                _text(case.get("execution_notes"), "No execution notes recorded."),
                "",
            ])
    warnings = _items((suite or {}).get("coverage_warnings"))
    if warnings:
        lines.extend(["## Coverage Warnings", "", *(f"- {warning}" for warning in warnings), ""])
    lines.extend([
        "## Important Limitation",
        "",
        "A generated test case is not an executed test. This report only records outcomes submitted by an authorized user; automated IDE/container execution is a future milestone.",
        "",
    ])
    return "\n".join(lines)


def render_traceability_csv(
    project: dict,
    artifacts: dict,
    suite: dict | None,
) -> str:
    requirements = _requirement_rows(artifacts.get("requirements") or {})
    plan = artifacts.get("project_plan") or {}
    architecture = artifacts.get("architecture") or {}
    review = artifacts.get("technical_review") or {}

    plan_refs: set[str] = set()
    for milestone in _items(plan.get("milestones")):
        plan_refs.update(str(ref).strip() for ref in _items(milestone.get("requirement_refs")))
        for task in _items(milestone.get("tasks")):
            plan_refs.update(str(ref).strip() for ref in _items(task.get("requirement_refs")))

    architecture_refs: set[str] = set()
    for item in _items(architecture.get("components")) + _items(architecture.get("api_endpoints")) + _items(architecture.get("architecture_decisions")):
        architecture_refs.update(str(ref).strip() for ref in _items(item.get("related_requirements")))

    review_coverage = {
        str(item.get("requirement_id", "")).strip(): item
        for item in _items(review.get("requirement_coverage"))
    }
    test_ids: dict[str, list[str]] = {}
    test_statuses: dict[str, list[str]] = {}
    for case in _items((suite or {}).get("cases")):
        requirement_id = str(case.get("requirement_id", "")).strip()
        test_ids.setdefault(requirement_id, []).append(str(case.get("test_id", "")))
        test_statuses.setdefault(requirement_id, []).append(str(case.get("status", "not_run")))

    buffer = io.StringIO()
    writer = csv.DictWriter(
        buffer,
        fieldnames=[
            "project_id",
            "project_name",
            "requirement_id",
            "requirement",
            "in_plan",
            "in_architecture",
            "review_status",
            "test_case_ids",
            "test_results",
        ],
    )
    writer.writeheader()
    project_name = _text((artifacts.get("requirements") or {}).get("project_name"), project.get("project_name") or "Specra Project")
    for requirement in requirements:
        requirement_id = requirement["requirement_id"]
        review_item = review_coverage.get(requirement_id, {})
        writer.writerow({
            "project_id": project.get("project_id", ""),
            "project_name": project_name,
            "requirement_id": requirement_id,
            "requirement": requirement["text"],
            "in_plan": "yes" if requirement_id in plan_refs else "no",
            "in_architecture": "yes" if requirement_id in architecture_refs else "no",
            "review_status": (
                "covered" if review_item and all(
                    review_item.get(flag, False)
                    for flag in ("in_project_plan", "in_architecture")
                ) else "gap" if review_item else "not_reviewed"
            ),
            "test_case_ids": "; ".join(test_ids.get(requirement_id, [])),
            "test_results": "; ".join(test_statuses.get(requirement_id, [])) or "no_test_cases",
        })
    return buffer.getvalue()


def safe_filename(project_name: str, project_id: str, suffix: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", project_name).strip("-").lower()[:48]
    slug = slug or "specra-project"
    short_id = re.sub(r"[^a-zA-Z0-9_-]", "", project_id)[:12]
    return f"{slug}-{short_id}.{suffix}"


def _markdown_inline(value: str) -> str:
    """Convert the small inline Markdown subset used by Specra into ReportLab markup."""
    import html

    escaped = html.escape(value, quote=False)
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"\x60([^\x60]+)\x60", r'<font name="Courier">\1</font>', escaped)
    return escaped


def _markdown_blocks(markdown: str) -> list[tuple[str, str]]:
    """Return (kind, text) blocks for the shared PDF and PowerPoint exporters."""
    blocks: list[tuple[str, str]] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(("paragraph", " ".join(part.strip() for part in paragraph).strip()))
            paragraph.clear()

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            flush()
        elif line.startswith("# "):
            flush()
            blocks.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            flush()
            blocks.append(("heading", line[3:].strip()))
        elif line.startswith("### "):
            flush()
            blocks.append(("subheading", line[4:].strip()))
        elif line.startswith("- "):
            flush()
            blocks.append(("bullet", line[2:].strip()))
        elif re.match(r"^\d+[.)]\s+", line):
            flush()
            blocks.append(("numbered", line))
        else:
            paragraph.append(line)
    flush()
    return blocks


def render_document_export(markdown: str, output_format: str = "md") -> tuple[bytes, str, str]:
    """Render a generated engineering document as Markdown, PDF, or PowerPoint."""
    fmt = output_format.strip().lower()
    if fmt in {"md", "markdown"}:
        return markdown.encode("utf-8"), "md", "text/markdown; charset=utf-8"

    blocks = _markdown_blocks(markdown)
    if fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

        output = io.BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=19 * mm,
            leftMargin=19 * mm,
            topMargin=19 * mm,
            bottomMargin=18 * mm,
            title=next((text for kind, text in blocks if kind == "title"), "Specra Engineering Document"),
            author="Specra",
        )
        base = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "SpecraTitle", parent=base["Title"], fontName="Helvetica-Bold",
            fontSize=21, leading=26, textColor=colors.HexColor("#17365D"),
            alignment=TA_LEFT, spaceAfter=12,
        )
        heading_style = ParagraphStyle(
            "SpecraHeading", parent=base["Heading1"], fontName="Helvetica-Bold",
            fontSize=14, leading=18, textColor=colors.HexColor("#245A81"),
            spaceBefore=12, spaceAfter=6, keepWithNext=True,
        )
        subheading_style = ParagraphStyle(
            "SpecraSubheading", parent=base["Heading2"], fontName="Helvetica-Bold",
            fontSize=11, leading=14, textColor=colors.HexColor("#354A5F"),
            spaceBefore=8, spaceAfter=4, keepWithNext=True,
        )
        body_style = ParagraphStyle(
            "SpecraBody", parent=base["BodyText"], fontName="Helvetica",
            fontSize=9.5, leading=13, spaceAfter=5, textColor=colors.HexColor("#202833"),
        )
        story = []
        for kind, text in blocks:
            safe_text = _markdown_inline(text)
            if kind == "title":
                story.extend([Paragraph(safe_text, title_style), HRFlowable(width="100%", color=colors.HexColor("#D7E3F0")), Spacer(1, 8)])
            elif kind == "heading":
                story.append(Paragraph(safe_text, heading_style))
            elif kind == "subheading":
                story.append(Paragraph(safe_text, subheading_style))
            elif kind == "bullet":
                story.append(Paragraph("&#8226;&nbsp; " + safe_text, body_style))
            else:
                story.append(Paragraph(safe_text, body_style))
        if not story:
            story.append(Paragraph("Specra Engineering Document", title_style))

        def add_page_number(canvas, doc) -> None:
            canvas.saveState()
            canvas.setStrokeColor(colors.HexColor("#D7E3F0"))
            canvas.line(19 * mm, 13 * mm, A4[0] - 19 * mm, 13 * mm)
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#64748B"))
            canvas.drawString(19 * mm, 8 * mm, "Generated by Specra - Review before approval")
            canvas.drawRightString(A4[0] - 19 * mm, 8 * mm, f"Page {doc.page}")
            canvas.restoreState()

        document.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
        return output.getvalue(), "pdf", "application/pdf"

    if fmt in {"ppt", "pptx", "powerpoint"}:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        presentation = Presentation()
        presentation.core_properties.title = next(
            (text for kind, text in blocks if kind == "title"),
            "Specra Engineering Document",
        )
        presentation.core_properties.author = "Specra"
        presentation.core_properties.subject = "Generated engineering document"
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        cover = presentation.slides.add_slide(presentation.slide_layouts[0])
        cover.shapes.title.text = presentation.core_properties.title
        cover.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
        cover.shapes.title.text_frame.paragraphs[0].font.bold = True
        cover.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(23, 54, 93)
        subtitle = cover.placeholders[1]
        subtitle.text = "Generated by Specra - Draft document - Review before approval"
        subtitle.text_frame.paragraphs[0].font.size = Pt(16)

        sections: list[tuple[str, list[tuple[str, str]]]] = []
        current_title = "Document Overview"
        current_items: list[tuple[str, str]] = []
        for kind, text in blocks:
            if kind == "title":
                continue
            if kind == "heading":
                if current_items:
                    sections.append((current_title, current_items))
                current_title = text
                current_items = []
                continue
            limit = 360
            if len(text) > limit:
                current_items.extend((kind, text[index:index + limit]) for index in range(0, len(text), limit))
            else:
                current_items.append((kind, text))
        if current_items:
            sections.append((current_title, current_items))
        if not sections:
            sections = [("Document Overview", [("paragraph", "No content was recorded.")])]

        for section_title, items in sections:
            chunks = [items[index:index + 6] for index in range(0, len(items), 6)] or [[]]
            for chunk_index, chunk in enumerate(chunks, start=1):
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = (
                    section_title if len(chunks) == 1
                    else f"{section_title} ({chunk_index}/{len(chunks)})"
                )
                slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
                slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(23, 54, 93)
                frame = slide.placeholders[1].text_frame
                frame.clear()
                frame.word_wrap = True
                for index, (kind, text) in enumerate(chunk):
                    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    paragraph.text = ("• " if kind == "bullet" else "") + text
                    paragraph.level = 0
                    paragraph.font.size = Pt(15 if kind in {"subheading", "heading"} else 13)
                    paragraph.font.bold = kind in {"subheading", "heading"}
                    paragraph.font.color.rgb = RGBColor(32, 40, 51)
                    paragraph.space_after = Pt(8)
                footer = slide.shapes.add_textbox(Inches(0.6), Inches(7.13), Inches(12.1), Inches(0.2))
                footer.text_frame.text = "Specra - Draft generated from saved project artifacts"
                footer.text_frame.paragraphs[0].font.size = Pt(8)
                footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(100, 116, 139)
                footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

        output = io.BytesIO()
        presentation.save(output)
        return (
            output.getvalue(),
            "pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

    raise ValueError("Unsupported document format. Choose md, pdf, or pptx.")
